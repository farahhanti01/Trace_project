from pathlib import Path
import re
import fitz  # PyMuPDF
from docx import Document
from openpyxl import load_workbook

from app.services.file_type_service import (
    IMAGE_EXTENSIONS,
    TRACE_TEXT_EXTENSIONS,
    is_image_extension,
    is_trace_extension,
)

SUPPORTED_TEXT_EXTENSIONS = TRACE_TEXT_EXTENSIONS
SUPPORTED_IMAGE_EXTENSIONS = IMAGE_EXTENSIONS

# Taille maximale d’un segment technique extrait.
# Ce ne sont pas encore les chunks du RAG.
MAX_SECTION_CHARACTERS = 250_000

FIELD_002_PATTERN = re.compile(
    r"(?P<prefix>\bFLD\s*\(?0*002\)?[^\[]*\[)"
    r"(?P<value>[^\]]+)"
    r"(?P<suffix>\])",
    flags=re.IGNORECASE,
)
IMAGE_FIELD_LINE_PATTERN = re.compile(
    r"\b(?:FLD|Field)\s*\(?0*(?P<number>\d{1,3})\)?"
    r"(?:\s*[:=\-]\s*|\s+)"
    r"(?:\(\s*\d+\s*\)\s*[:=\-]?\s*)?"
    r"(?:\[\s*)?(?P<value>[A-Za-z0-9*._\-/ ]{1,160})",
    flags=re.IGNORECASE,
)
IMAGE_MTI_PATTERN = re.compile(
    r"\bM\.?\s*T\.?\s*I\.?\b\s*[:=\-]?\s*\[?\s*(?P<mti>\d{4})\s*\]?"
    r"|\bMTI\b\s*[:=\-]?\s*\[?\s*(?P<mti_alt>\d{4})\s*\]?",
    flags=re.IGNORECASE,
)


class UnsupportedDocumentTypeError(ValueError):
    """Raised when the document type is not supported."""


class DocumentExtractionError(RuntimeError):
    """Raised when document extraction fails."""


def read_text_with_fallback(
    file_path: Path,
) -> tuple[str, str]:
    """
    Read a text file by trying the most common encodings.

    Returns:
        A tuple containing:
        - extracted text
        - detected encoding
    """

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1",
    ]

    last_error = None

    for encoding in encodings:
        try:
            text = file_path.read_text(
                encoding=encoding,
            )

            return text, encoding

        except UnicodeDecodeError as error:
            last_error = error

    raise DocumentExtractionError(
        f"Unable to decode the file '{file_path.name}'."
    ) from last_error


def normalize_text(text: str) -> str:
    """
    Apply minimal cleaning without modifying
    the technical meaning of logs or specifications.
    """

    normalized = text.replace(
        "\r\n",
        "\n",
    ).replace(
        "\r",
        "\n",
    )

    # Remove null characters sometimes found in exports.
    normalized = normalized.replace(
        "\x00",
        "",
    )

    return normalized.strip()


def mask_field_002_value(value: str) -> str:
    digits = re.sub(r"\D", "", value)

    if len(digits) <= 10:
        return "******"

    return f"{digits[:6]}******{digits[-4:]}"


def mask_iso_field_002(text: str) -> str:
    return FIELD_002_PATTERN.sub(
        lambda match: (
            f"{match.group('prefix')}"
            f"{mask_field_002_value(match.group('value'))}"
            f"{match.group('suffix')}"
        ),
        text,
    )


def optional_ocr_image_text(file_path: Path) -> tuple[str, str]:
    """
    Extract text from an image when optional OCR dependencies are installed.

    This keeps screenshot ingestion best-effort: the document remains accepted
    even on machines without Pillow, pytesseract, or the Tesseract binary.
    """

    try:
        from PIL import Image
    except ImportError:
        return "", "image:no_ocr_pillow_missing"

    try:
        import pytesseract
    except ImportError:
        return "", "image:no_ocr_pytesseract_missing"

    try:
        with Image.open(file_path) as image:
            text = pytesseract.image_to_string(image)
    except Exception as error:
        return "", f"image:ocr_failed:{type(error).__name__}"

    return normalize_text(text), "image:ocr"


def normalize_visible_field_number(value: str) -> str:
    digits = re.sub(r"\D", "", value)

    if not digits:
        return value.strip()

    return f"{int(digits):03d}"


def build_image_overview_text(
    *,
    filename: str,
    ocr_text: str,
    encoding: str,
) -> str:
    if ocr_text:
        return (
            f"Capture d'ecran importee: {filename}\n"
            "Texte visible extrait par OCR. Cette capture peut etre "
            "utilisee comme source documentaire pour decrire les elements "
            "visibles et poser des questions contextuelles sur une trace, "
            "des Fields, des fonctions, des erreurs ou des echanges HSM."
        )

    return (
        f"Capture d'ecran importee: {filename}\n"
        "Aucun texte OCR exploitable n'a ete extrait de cette image. "
        "Le fichier est conserve comme source, mais l'analyse documentaire "
        "necessite l'installation d'un moteur OCR local pour lire les "
        f"elements visibles. Statut OCR: {encoding}."
    )


def image_field_sections(ocr_text: str) -> list[dict]:
    sections = []

    for line_number, line in enumerate(ocr_text.splitlines(), start=1):
        match = IMAGE_FIELD_LINE_PATTERN.search(line)

        if not match:
            continue

        field_number = normalize_visible_field_number(match.group("number"))
        value = match.group("value").strip(" []:=-")

        if field_number == "002":
            value = mask_field_002_value(value)

        sections.append({
            "text": (
                f"Element visible dans la capture: Field {field_number} "
                f"avec la valeur {value}.\n"
                f"Ligne OCR source: {line.strip()}"
            ),
            "page": None,
            "paragraph": line_number,
            "heading": f"Capture d'ecran OCR > Field {field_number}",
            "field_number": field_number,
            "content_type": "field_related",
        })

    return sections


def image_mti_sections(ocr_text: str) -> list[dict]:
    sections = []

    for line_number, line in enumerate(ocr_text.splitlines(), start=1):
        match = IMAGE_MTI_PATTERN.search(line)

        if not match:
            continue

        mti = match.group("mti") or match.group("mti_alt")

        sections.append({
            "text": (
                f"Element visible dans la capture: MTI {mti}.\n"
                f"Ligne OCR source: {line.strip()}"
            ),
            "page": None,
            "paragraph": line_number,
            "heading": "Capture d'ecran OCR > MTI",
            "content_type": "message_type",
        })

    return sections


def image_hsm_sections(ocr_text: str) -> list[dict]:
    hsm_lines = [
        line.strip()
        for line in ocr_text.splitlines()
        if re.search(
            r"\b(HSM|TO HSM|FROM HSM|HsmResultCode|command_[A-Z0-9]+)\b",
            line,
            flags=re.IGNORECASE,
        )
    ]

    if not hsm_lines:
        return []

    return [{
        "text": (
            "Elements HSM visibles dans la capture:\n"
            + "\n".join(hsm_lines)
        ),
        "page": None,
        "paragraph": None,
        "heading": "Capture d'ecran OCR > HSM",
        "content_type": "description",
    }]


def image_function_sections(ocr_text: str) -> list[dict]:
    function_lines = [
        line.strip()
        for line in ocr_text.splitlines()
        if re.search(
            r"\b(Start|End)\s+[A-Za-z_][A-Za-z0-9_]*\b"
            r"|\b(ERROR|FAILED|SUCCESS|NOK)\b",
            line,
            flags=re.IGNORECASE,
        )
    ]

    if not function_lines:
        return []

    return [{
        "text": (
            "Fonctions, statuts ou erreurs visibles dans la capture:\n"
            + "\n".join(function_lines)
        ),
        "page": None,
        "paragraph": None,
        "heading": "Capture d'ecran OCR > Fonctions et erreurs",
        "content_type": "description",
    }]


def build_image_sections(
    *,
    file_path: Path,
    ocr_text: str,
    encoding: str,
) -> list[dict]:
    sections = [{
        "text": build_image_overview_text(
            filename=file_path.name,
            ocr_text=ocr_text,
            encoding=encoding,
        ),
        "page": None,
        "paragraph": 1,
        "heading": "Capture d'ecran OCR > Vue generale",
        "content_type": "description",
    }]

    if ocr_text:
        masked_text = mask_iso_field_002(ocr_text)
        sections.append({
            "text": masked_text,
            "page": None,
            "paragraph": 2,
            "heading": "Capture d'ecran OCR > Texte extrait",
            "content_type": "description",
        })
        sections.extend(image_mti_sections(masked_text))
        sections.extend(image_field_sections(masked_text))
        sections.extend(image_hsm_sections(masked_text))
        sections.extend(image_function_sections(masked_text))

    return sections


def split_paragraphs(text: str) -> list[str]:
    """
    Split extracted text into readable paragraphs while preserving
    technical lines when the source has no blank-line structure.
    """

    cleaned = normalize_text(text)

    if not cleaned:
        return []

    if "\n\n" in cleaned:
        return [
            paragraph.strip()
            for paragraph in cleaned.split("\n\n")
            if paragraph.strip()
        ]

    return [
        line.strip()
        for line in cleaned.splitlines()
        if line.strip()
    ]


def split_extracted_text(
    text: str,
    max_characters: int = MAX_SECTION_CHARACTERS,
) -> list[str]:
    """
    Split a large extracted text into technical sections.

    The function tries to preserve line boundaries.
    These sections are not the semantic chunks used by the RAG.
    """

    if not text:
        return []

    sections = []
    current_lines = []
    current_size = 0

    for line in text.splitlines(keepends=True):
        line_size = len(line)

        if (
            current_lines
            and current_size + line_size > max_characters
        ):
            section = "".join(current_lines).strip()

            if section:
                sections.append(section)

            current_lines = []
            current_size = 0

        current_lines.append(line)
        current_size += line_size

    if current_lines:
        section = "".join(current_lines).strip()

        if section:
            sections.append(section)

    return sections


def build_pdf_page_headings(
    pdf: fitz.Document,
) -> dict[int, str]:
    """
    Build an approximate page -> heading map from the PDF outline.
    This gives the RAG chunks useful chapter/section context.
    """

    toc = pdf.get_toc()

    if not toc:
        return {}

    headings_by_page = {}
    active_headings: dict[int, str] = {}
    toc_index = 0

    for page_number in range(1, pdf.page_count + 1):
        while (
            toc_index < len(toc)
            and toc[toc_index][2] <= page_number
        ):
            level, title, _ = toc[toc_index]
            active_headings[level] = normalize_text(title)

            for stale_level in list(active_headings):
                if stale_level > level:
                    del active_headings[stale_level]

            toc_index += 1

        if active_headings:
            headings_by_page[page_number] = " > ".join(
                active_headings[level]
                for level in sorted(active_headings)
            )

    return headings_by_page


def detect_pdf_page_heading(
    text: str,
    fallback_heading: str | None,
) -> str | None:
    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    for index, line in enumerate(lines[:-1]):
        if not re_match_section_number(line):
            continue

        next_line = lines[index + 1]

        if len(next_line) > 80:
            continue

        detected_heading = f"{line} {next_line}"

        if fallback_heading:
            detected_number = line.strip()
            parents = []

            for segment in fallback_heading.split(" > "):
                match = re.match(
                    r"^(\d+(?:\.\d+)*)\b",
                    segment,
                )

                if segment.startswith("Chapter "):
                    parents.append(segment)
                    continue

                if (
                    match
                    and detected_number.startswith(match.group(1))
                    and match.group(1) != detected_number
                ):
                    parents.append(segment)

            return " > ".join([
                *parents,
                detected_heading,
            ])

        return detected_heading

    return fallback_heading


def re_match_section_number(value: str) -> bool:
    return bool(
        re.match(
            r"^\d+(?:\.\d+){1,5}$",
            value.strip(),
        )
    )


def extract_text_file(
    file_path: Path,
) -> dict:
    """
    Extract content from TXT and LOG files.
    """

    text, encoding = read_text_with_fallback(
        file_path
    )

    cleaned_text = mask_iso_field_002(
        normalize_text(text)
    )

    sections = split_extracted_text(
        cleaned_text
    )

    return {
        "text": cleaned_text,
        "sections": sections,
        "encoding": encoding,
        "content_length": len(cleaned_text),
        "line_count": (
            cleaned_text.count("\n") + 1
            if cleaned_text
            else 0
        ),
    }

def extract_pdf(file_path: Path) -> dict:
    pdf = fitz.open(file_path)
    headings_by_page = build_pdf_page_headings(pdf)

    pages = []
    full_text = []
    sections = []

    for page_number, page in enumerate(pdf, start=1):
        text = normalize_text(page.get_text())
        page_heading = detect_pdf_page_heading(
            text,
            headings_by_page.get(page_number),
        )

        if text:
            pages.append({
                "page": page_number,
                "text": text,
            })

            full_text.append(text)

            sections.append({
                "text": text,
                "page": page_number,
                "paragraph": None,
                "heading": page_heading,
            })

    merged = "\n\n".join(full_text)

    return {
        "text": merged,
        "sections": sections,
        "pages": pages,
        "encoding": "pdf",
        "content_length": len(merged),
        "line_count": merged.count("\n") + 1,
    }

def extract_docx(file_path: Path) -> dict:
    document = Document(file_path)

    paragraphs = []
    sections = []

    for paragraph_number, paragraph in enumerate(
        document.paragraphs,
        start=1,
    ):
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)
            sections.append({
                "text": text,
                "paragraph": paragraph_number,
                "heading": None,
            })

    merged = "\n".join(paragraphs)

    return {
        "text": merged,
        "sections": sections,
        "encoding": "docx",
        "content_length": len(merged),
        "line_count": merged.count("\n") + 1,
    }

def extract_xlsx(file_path: Path) -> dict:
    workbook = load_workbook(
        filename=file_path,
        data_only=True,
    )

    sheets = []

    merged = []

    for sheet in workbook.worksheets:

        rows = []
        section_rows = []

        for row_number, row in enumerate(
            sheet.iter_rows(values_only=True),
            start=1,
        ):

            cells = []

            for column_number, value in enumerate(row, start=1):
                if value is None:
                    continue

                cleaned_value = str(value).strip()

                if not cleaned_value:
                    continue

                cells.append({
                    "column": column_number,
                    "value": cleaned_value,
                })

            values = [
                cell["value"]
                for cell in cells
            ]

            if values:
                row_text = " | ".join(values)
                rows.append(row_text)
                section_rows.append({
                    "text": row_text,
                    "sheet": sheet.title,
                    "paragraph": row_number,
                    "row_number": row_number,
                    "heading": sheet.title,
                    "cells": cells,
                })

        text = "\n".join(rows)

        if text:
            sheets.append({
                "sheet": sheet.title,
                "text": text,
                "sections": section_rows,
            })

            merged.append(text)

    final_text = "\n\n".join(merged)

    return {
        "text": final_text,
        "sections": [
            section
            for sheet in sheets
            for section in sheet["sections"]
        ],
        "encoding": "xlsx",
        "content_length": len(final_text),
        "line_count": final_text.count("\n") + 1,
    }


def extract_image(file_path: Path) -> dict:
    """
    Extract visible text from a screenshot/image when OCR is available.

    The returned sections are intentionally text-first so the existing RAG
    pipeline can index screenshots without a separate vision pipeline.
    """

    ocr_text, encoding = optional_ocr_image_text(file_path)
    cleaned_text = mask_iso_field_002(
        normalize_text(ocr_text)
    )
    sections = build_image_sections(
        file_path=file_path,
        ocr_text=cleaned_text,
        encoding=encoding,
    )
    merged = "\n\n".join(
        section["text"]
        for section in sections
        if section.get("text")
    )

    return {
        "text": merged,
        "sections": sections,
        "encoding": encoding,
        "content_length": len(merged),
        "line_count": (
            merged.count("\n") + 1
            if merged
            else 0
        ),
    }

def extract_document(
    file_path: Path,
) -> dict:
    """
    Select the correct extractor according
    to the document extension.
    """

    extension = file_path.suffix.lower()

    if is_trace_extension(extension):
        return extract_text_file(file_path)

    if is_image_extension(extension):
        return extract_image(file_path)

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".xlsx":
        return extract_xlsx(file_path)

    raise UnsupportedDocumentTypeError(
        f"Unsupported extension: {extension}"
    )

    raise UnsupportedDocumentTypeError(
        f"Extraction is not yet supported for "
        f"the extension '{extension}'."
    )
