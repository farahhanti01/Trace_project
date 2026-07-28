from pathlib import Path
from typing import Any
from html import escape

import fitz
from bson import ObjectId
from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse, HTMLResponse, Response
from openpyxl import load_workbook

from app.database import (
    conversations_collection,
    document_sections_collection,
    documents_collection,
)
from app.services.document_service import (
    BACKEND_ROOT,
    get_document,
    # reindex_document,
    serialize_document,
)
from app.services.log_analysis_agent_service import (
    apply_deterministic_enrichment,
    enrich_transactions_with_documentation,
    load_reference_sections,
    response_transactions,
)
from app.services.log_parser_service import (
    build_statistics,
    parse_log_transactions,
)


router = APIRouter(
    prefix="/api/admin",
    tags=["Admin"],
)


LOG_EXTENSIONS = {".txt", ".log"}
INLINE_MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".txt": "text/plain; charset=utf-8",
    ".log": "text/plain; charset=utf-8",
    ".docx": (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document"
    ),
    ".xlsx": (
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    ),
}


def require_object_id(
    value: str,
    label: str = "ID",
) -> ObjectId:
    if not ObjectId.is_valid(value):
        raise HTTPException(
            status_code=400,
            detail=f"Invalid {label}.",
        )

    return ObjectId(value)


async def count_sections_by_document_id(
    document_ids: list[str],
) -> dict[str, int]:
    if not document_ids:
        return {}

    pipeline = [
        {
            "$match": {
                "document_id": {"$in": document_ids},
            }
        },
        {
            "$group": {
                "_id": "$document_id",
                "count": {"$sum": 1},
            }
        },
    ]

    rows = await document_sections_collection.aggregate(
        pipeline
    ).to_list(length=1_000)

    return {
        row["_id"]: row["count"]
        for row in rows
    }


async def list_serialized_documents(
    query: dict[str, Any] | None = None,
    limit: int = 1_000,
) -> list[dict[str, Any]]:
    documents = await documents_collection.find(
        query or {}
    ).sort(
        "created_at",
        -1,
    ).to_list(length=limit)

    section_counts = await count_sections_by_document_id(
        [str(document["_id"]) for document in documents]
    )

    serialized = []

    for document in documents:
        item = serialize_document(document)
        item["section_count"] = section_counts.get(
            str(document["_id"]),
            item.get("section_count", 0),
        )
        serialized.append(item)

    return serialized


def document_kind(
    document: dict[str, Any],
) -> str:
    extension = document.get("extension")

    if extension in LOG_EXTENSIONS:
        return "trace"

    if extension == ".xlsx":
        return "excel"

    if extension == ".pdf":
        return "pdf"

    return "document"


def summarize_block(
    conversation: dict[str, Any],
    documents: list[dict[str, Any]],
) -> dict[str, Any]:
    traces = [
        document
        for document in documents
        if document.get("extension") in LOG_EXTENSIONS
    ]
    references = [
        document
        for document in documents
        if document.get("extension") not in LOG_EXTENSIONS
    ]
    failed_documents = [
        document
        for document in documents
        if str(document.get("status", "")).endswith("failed")
    ]

    return {
        "id": str(conversation["_id"]),
        "conversation_id": str(conversation["_id"]),
        "title": conversation.get("title", "Untitled"),
        "agent": conversation.get("agent"),
        "created_at": conversation.get("created_at"),
        "updated_at": conversation.get("updated_at"),
        "status": "failed" if failed_documents else "ready",
        "trace_count": len(traces),
        "reference_count": len(references),
        "documents_count": len(documents),
        "traces": traces,
        "references": references,
        "documents": documents,
    }


@router.get("/documents")
async def admin_documents(
    type: str | None = None,
    agent: str | None = None,
    status: str | None = None,
) -> list[dict[str, Any]]:
    query: dict[str, Any] = {}

    if agent:
        query["agent"] = agent

    if status:
        query["status"] = status

    if type:
        normalized_type = type.lower()

        if normalized_type == "trace":
            query["extension"] = {"$in": list(LOG_EXTENSIONS)}
        elif normalized_type in {"pdf", "xlsx", "docx"}:
            query["extension"] = f".{normalized_type}"

    documents = await list_serialized_documents(query)

    for document in documents:
        document["kind"] = document_kind(document)

    return documents


@router.get("/documents/{document_id}")
async def admin_document_detail(
    document_id: str,
) -> dict[str, Any]:
    document = await get_document(document_id)

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    sections = await document_sections_collection.find(
        {
            "document_id": document_id,
        }
    ).sort(
        [
            ("section_index", 1),
            ("chunk_index", 1),
        ]
    ).to_list(length=5_000)

    preview_sections = [
        {
            "section_index": section.get("section_index"),
            "chunk_index": section.get("chunk_index"),
            "page": section.get("page"),
            "sheet": section.get("sheet"),
            "paragraph": section.get("paragraph"),
            "heading": section.get("heading"),
            "character_count": section.get("character_count"),
            "embedding_status": (
                "embedded"
                if section.get("embedding")
                else "missing"
            ),
            "text": section.get("text", ""),
        }
        for section in sections
    ]

    document["kind"] = document_kind(document)
    document["sections"] = preview_sections

    return document


@router.get("/documents/{document_id}/download")
async def admin_download_document(
    document_id: str,
):
    document = await get_document(document_id)

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    file_path = BACKEND_ROOT / document["relative_path"]

    if not file_path.exists():
        raise HTTPException(
            status_code=404,
            detail="Stored file not found.",
        )

    return FileResponse(
        path=file_path,
        filename=document["original_filename"],
        media_type=document.get("content_type")
        or "application/octet-stream",
    )


@router.get("/documents/{document_id}/view")
async def admin_view_document(
    document_id: str,
):
    document = await get_document(document_id)

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    file_path = BACKEND_ROOT / document["relative_path"]

    if not file_path.exists():
        raise HTTPException(
            status_code=404,
            detail="Stored file not found.",
        )

    response = FileResponse(
        path=file_path,
        filename=document["original_filename"],
        media_type=INLINE_MEDIA_TYPES.get(
            document.get("extension"),
            document.get("content_type") or "application/octet-stream",
        ),
    )
    response.headers["Content-Disposition"] = (
        f'inline; filename="{document["original_filename"]}"'
    )

    return response


@router.get("/documents/{document_id}/preview")
async def admin_preview_document(
    document_id: str,
):
    document = await get_document(document_id)

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    if document.get("extension") == ".pdf":
        return await admin_view_document(document_id)

    extension = document.get("extension")
    file_path = BACKEND_ROOT / document["relative_path"]

    if extension == ".xlsx":
        if not file_path.exists():
            raise HTTPException(
                status_code=404,
                detail="Stored file not found.",
            )

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )
        sheets_html = []

        for worksheet in workbook.worksheets:
            rows_html = []

            for row_number, row in enumerate(
                worksheet.iter_rows(),
                start=1,
            ):
                cells_html = []
                has_value = False

                for cell in row:
                    value = getattr(cell, "value", None)
                    tag = "th" if row_number == 1 else "td"

                    if value is not None:
                        has_value = True

                    cells_html.append(
                        f"<{tag}>{escape(str(value or ''))}</{tag}>"
                    )

                if has_value:
                    rows_html.append(f"<tr>{''.join(cells_html)}</tr>")

            sheets_html.append(
                "<section>"
                f"<h2>{escape(worksheet.title)}</h2>"
                "<div class=\"sheet-scroll\">"
                f"<table>{''.join(rows_html)}</table>"
                "</div>"
                "</section>"
            )

        html = f"""
        <!doctype html>
        <html>
          <head>
            <title>{escape(document["original_filename"])}</title>
            <style>
              body {{
                margin: 24px;
                color: #172033;
                background: #ffffff;
                font-family: Arial, sans-serif;
                font-size: 13px;
              }}
              h1 {{
                margin: 0 0 6px;
                font-size: 22px;
              }}
              .sub {{
                margin: 0 0 18px;
                color: #667085;
              }}
              section {{
                margin: 0 0 22px;
              }}
              h2 {{
                margin: 0 0 8px;
                font-size: 15px;
              }}
              .sheet-scroll {{
                max-width: 100%;
                overflow: auto;
                border: 1px solid #d9e0ea;
                border-radius: 8px;
              }}
              table {{
                border-collapse: collapse;
                min-width: 100%;
                background: #ffffff;
              }}
              th,
              td {{
                padding: 7px 9px;
                border: 1px solid #e4e7ec;
                text-align: left;
                vertical-align: top;
                white-space: pre-wrap;
                min-width: 120px;
              }}
              th {{
                position: sticky;
                top: 0;
                z-index: 1;
                background: #f2f4f7;
                font-weight: 800;
              }}
            </style>
          </head>
          <body>
            <h1>{escape(document["original_filename"])}</h1>
            <p class="sub">Workbook preview</p>
            {''.join(sheets_html) or "<p>No workbook content available.</p>"}
          </body>
        </html>
        """

        workbook.close()
        return HTMLResponse(html)

    if extension in LOG_EXTENSIONS:
        preview_text = read_stored_text_document(document)
        subtitle = "Stored file preview"
    else:
        sections = await document_sections_collection.find(
            {
                "document_id": document_id,
            }
        ).sort(
            [
                ("section_index", 1),
                ("chunk_index", 1),
            ]
        ).to_list(length=5_000)
        preview_text = "\n\n".join(
            section.get("text", "")
            for section in sections
            if section.get("text")
        )
        subtitle = "Extracted document preview"

    content_html = (
        f"<pre>{escape(preview_text)}</pre>"
        if preview_text
        else "<p>No preview content available.</p>"
    )

    html = f"""
    <!doctype html>
    <html>
      <head>
        <title>{escape(document["original_filename"])}</title>
        <style>
          body {{
            margin: 32px;
            color: #172033;
            background: #ffffff;
            font-family: Arial, sans-serif;
            font-size: 13px;
            line-height: 1.5;
          }}
          h1 {{
            margin: 0 0 6px;
            font-size: 22px;
          }}
          .sub {{
            margin: 0 0 18px;
            color: #667085;
          }}
          pre {{
            margin: 0;
            padding: 16px;
            border: 1px solid #d9e0ea;
            border-radius: 8px;
            background: #ffffff;
            white-space: pre;
            overflow: auto;
            font-family: Consolas, monospace;
            font-size: 12px;
            line-height: 1.55;
          }}
        </style>
      </head>
      <body>
        <h1>{escape(document["original_filename"])}</h1>
        <p class="sub">{subtitle}</p>
        {content_html}
      </body>
    </html>
    """

    return HTMLResponse(html)


# @router.post("/documents/{document_id}/reindex")
# async def admin_reindex_document(
#     document_id: str,
# ):
#     return await reindex_document(document_id)


@router.delete("/documents/{document_id}")
async def admin_delete_document(
    document_id: str,
) -> dict[str, str]:
    object_id = require_object_id(document_id, "document ID")
    document = await documents_collection.find_one({"_id": object_id})

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    file_path = BACKEND_ROOT / document["relative_path"]

    await documents_collection.delete_one({"_id": object_id})
    await document_sections_collection.delete_many(
        {
            "document_id": document_id,
        }
    )

    if file_path.exists():
        file_path.unlink()

    return {
        "status": "deleted",
        "document_id": document_id,
    }


@router.get("/analysis-blocks")
async def admin_analysis_blocks() -> list[dict[str, Any]]:
    conversations = await conversations_collection.find({}).sort(
        "updated_at",
        -1,
    ).to_list(length=500)
    documents = await list_serialized_documents(limit=2_000)
    documents_by_conversation: dict[str, list[dict[str, Any]]] = {}

    for document in documents:
        document["kind"] = document_kind(document)
        documents_by_conversation.setdefault(
            document["conversation_id"],
            [],
        ).append(document)

    blocks = []

    for conversation in conversations:
        conversation_id = str(conversation["_id"])
        conversation_documents = documents_by_conversation.get(
            conversation_id,
            [],
        )

        if not conversation_documents:
            continue

        blocks.append(
            summarize_block(
                conversation=conversation,
                documents=conversation_documents,
            )
        )

    return blocks


def read_stored_text_document(
    document: dict[str, Any],
) -> str:
    file_path = BACKEND_ROOT / document["relative_path"]

    if not file_path.exists():
        raise HTTPException(
            status_code=404,
            detail="Stored log file not found.",
        )

    try:
        return file_path.read_text(
            encoding=document.get("encoding") or "utf-8",
        )
    except UnicodeDecodeError:
        return file_path.read_text(encoding="latin-1")


async def build_admin_log_story(
    document_id: str,
    failed_only: bool = False,
) -> dict[str, Any]:
    document = await get_document(document_id)

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    if document.get("extension") not in LOG_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Log Story is available only for .txt/.log files.",
        )

    text = read_stored_text_document(document)
    transactions = parse_log_transactions(
        text=text,
        source=document["original_filename"],
    )
    reference_sections = await load_reference_sections(
        conversation_id=document["conversation_id"],
        referenced_document_ids=[],
    )

    await enrich_transactions_with_documentation(
        transactions=transactions,
        reference_sections=reference_sections,
    )

    statistics = build_statistics(transactions)
    apply_deterministic_enrichment(transactions)

    if failed_only:
        selected_transactions = response_transactions(
            transactions=transactions,
            question="failed error nok echec",
        )
    else:
        selected_transactions = transactions

    return {
        "source": document["original_filename"],
        "document_id": document_id,
        "conversation_id": document["conversation_id"],
        "summary": (
            f"{statistics['total_transactions']} transaction(s) detected: "
            f"{statistics['successful_transactions']} success, "
            f"{statistics['failed_transactions']} failed, "
            f"{statistics['warning_transactions']} warning."
        ),
        "statistics": {
            **statistics,
            "returned_transactions": len(selected_transactions),
        },
        "transactions": selected_transactions,
    }


def draw_text(
    page: fitz.Page,
    text: Any,
    rect: fitz.Rect,
    *,
    size: float = 9,
    bold: bool = False,
    color: tuple[float, float, float] = (0, 0, 0),
) -> None:
    page.insert_textbox(
        rect,
        str(text or ""),
        fontsize=size,
        fontname="hebo" if bold else "helv",
        color=color,
        align=fitz.TEXT_ALIGN_LEFT,
    )


def build_log_story_pdf(
    log_story: dict[str, Any],
) -> bytes:
    pdf = fitz.open()
    page_width = 595
    page_height = 842
    margin = 54
    content_width = page_width - (margin * 2)
    y = margin

    def new_page() -> fitz.Page:
        return pdf.new_page(width=page_width, height=page_height)

    page = new_page()

    def ensure_space(required_height: float) -> None:
        nonlocal page, y

        if y + required_height <= page_height - margin:
            return

        page = new_page()
        y = margin

    draw_text(
        page,
        f"Log Story - {log_story.get('source', 'trace')}",
        fitz.Rect(margin, y, page_width - margin, y + 24),
        size=16,
        bold=True,
    )
    y += 30
    draw_text(
        page,
        log_story.get("summary", ""),
        fitz.Rect(margin, y, page_width - margin, y + 18),
        size=9,
    )
    y += 28

    for transaction in log_story.get("transactions", []):
        fields = transaction.get("fields") or {}
        story = transaction.get("log_story") or []
        hsm_commands = (
            (transaction.get("hsm_analysis") or {}).get("commands") or []
        )
        row_height = 18
        hsm_height = 0

        if hsm_commands:
            hsm_height = 34 + (len(hsm_commands) * row_height)

        card_height = 86 + hsm_height + (max(len(story), 1) * row_height)

        ensure_space(card_height + 16)

        card = fitz.Rect(
            margin,
            y,
            margin + content_width,
            y + card_height,
        )
        page.draw_rect(
            card,
            color=(0.84, 0.88, 0.94),
            fill=None,
            width=0.7,
        )

        draw_text(
            page,
            transaction.get("display_name") or transaction.get("transaction_id"),
            fitz.Rect(margin + 10, y + 10, margin + content_width - 90, y + 28),
            size=11,
            bold=True,
        )
        draw_text(
            page,
            transaction.get("status", ""),
            fitz.Rect(margin + content_width - 70, y + 10, margin + content_width - 10, y + 28),
            size=9,
            bold=True,
        )

        field_labels = [
            ("MTI", transaction.get("mti") or "UNKNOWN"),
            ("FLD 002", fields.get("002") or "N/A"),
            ("FLD 003", fields.get("003") or "N/A"),
            ("FLD 037", fields.get("037") or "N/A"),
            ("FLD 039", fields.get("039") or "N/A"),
        ]
        field_y = y + 38
        field_gap = 6
        field_width = (content_width - 20 - (field_gap * 4)) / 5

        for index, (label, value) in enumerate(field_labels):
            x = margin + 10 + (index * (field_width + field_gap))
            box = fitz.Rect(x, field_y, x + field_width, field_y + 34)
            page.draw_rect(
                box,
                color=(0.88, 0.91, 0.96),
                fill=(0.99, 1, 1),
                width=0.5,
            )
            draw_text(
                page,
                label,
                fitz.Rect(x + 5, field_y + 4, x + field_width - 5, field_y + 12),
                size=6,
                bold=True,
            )
            draw_text(
                page,
                value,
                fitz.Rect(x + 5, field_y + 15, x + field_width - 5, field_y + 30),
                size=7,
            )

        table_y = y + 84

        if hsm_commands:
            draw_text(
                page,
                "HSM Analysis",
                fitz.Rect(margin + 12, table_y, margin + 150, table_y + 14),
                size=8,
                bold=True,
            )
            hsm_thread = (transaction.get("hsm_analysis") or {}).get("thread")
            if hsm_thread:
                draw_text(
                    page,
                    f"Thread {hsm_thread}",
                    fitz.Rect(margin + content_width - 150, table_y, margin + content_width - 10, table_y + 14),
                    size=8,
                    bold=True,
                )
            table_y += row_height

            for command in hsm_commands:
                draw_text(
                    page,
                    f"{command.get('command') or 'N/A'} -> {command.get('response_command') or 'N/A'}",
                    fitz.Rect(margin + 12, table_y, margin + 160, table_y + 14),
                    size=8,
                )
                draw_text(
                    page,
                    f"HsmResultCode {command.get('hsm_result_code') or 'N/A'}",
                    fitz.Rect(margin + 170, table_y, margin + 330, table_y + 14),
                    size=8,
                )
                draw_text(
                    page,
                    command.get("status") or "UNKNOWN",
                    fitz.Rect(margin + content_width - 150, table_y, margin + content_width - 10, table_y + 14),
                    size=8,
                    bold=True,
                )
                table_y += row_height

        draw_text(
            page,
            "#",
            fitz.Rect(margin + 12, table_y, margin + 40, table_y + 14),
            size=8,
            bold=True,
        )
        draw_text(
            page,
            "Function",
            fitz.Rect(margin + 58, table_y, margin + 250, table_y + 14),
            size=8,
            bold=True,
        )
        draw_text(
            page,
            "Status",
            fitz.Rect(margin + content_width - 150, table_y, margin + content_width - 10, table_y + 14),
            size=8,
            bold=True,
        )
        table_y += row_height

        for item in story:
            page.draw_line(
                fitz.Point(margin + 10, table_y - 4),
                fitz.Point(margin + content_width - 10, table_y - 4),
                color=(0.88, 0.91, 0.96),
                width=0.5,
            )
            draw_text(
                page,
                item.get("order"),
                fitz.Rect(margin + 12, table_y, margin + 40, table_y + 14),
                size=8,
            )
            draw_text(
                page,
                item.get("function_name"),
                fitz.Rect(margin + 58, table_y, margin + 330, table_y + 14),
                size=8,
            )
            draw_text(
                page,
                item.get("status"),
                fitz.Rect(margin + content_width - 150, table_y, margin + content_width - 10, table_y + 14),
                size=8,
                bold=True,
            )
            table_y += row_height

        y += card_height + 14

    return pdf.tobytes(
        garbage=4,
        deflate=True,
    )


@router.get("/documents/{document_id}/log-story")
async def admin_document_log_story(
    document_id: str,
    failed_only: bool = False,
) -> dict[str, Any]:
    return await build_admin_log_story(
        document_id=document_id,
        failed_only=failed_only,
    )


@router.get("/documents/{document_id}/log-story/pdf")
async def admin_document_log_story_pdf(
    document_id: str,
    failed_only: bool = False,
):
    log_story = await build_admin_log_story(
        document_id=document_id,
        failed_only=failed_only,
    )
    pdf_bytes = build_log_story_pdf(log_story)
    filename = f"log-story-{Path(log_story.get('source', 'trace')).stem}.pdf"

    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'inline; filename="{filename}"',
        },
    )
